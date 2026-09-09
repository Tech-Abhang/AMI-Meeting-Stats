#!/usr/bin/env python3
"""
create_negative_trials.py

Builds the capped NEGATIVE trial table (different speaker, same gender,
different meeting) using the 50 equal-width duration_diff bins already
computed from the real positive trials, at a configurable ratio (default
10x) per bin -- exactly the scheme worked out in conversation with your
professor.

OUTPUT: a single pandas DataFrame, saved with joblib, 3 columns:
    enroll_path , test_path , duration_diff
identical shape to positive_trials.csv, so the two can be told apart
only by which file they came from (or a label you add yourself later).
Path columns are stored as pandas 'category' dtype -- when you load
this and call df.head(), you see real, full, readable path text. The
compactness (same idea as the earlier parquet approach) happens
automatically underneath, invisibly.

HOW BINS AND TARGETS ARE DECIDED
---------------------------------
Reads positive_trial_full_distribution_stats.json (already generated
from your real positive_trials.csv) and uses its "equal_width_bins"
table directly -- the real 50 bin edges and real positive counts, not
re-derived here. For each bin:
    target_for_this_bin = positive_count_in_bin * ratio
Example: bin 0 (0.0s-0.66s) has 7,079,102 real positive trials, so at
ratio=10 the target for that bin is 70,791,020 negatives.

If a bin's negative candidate pool has FEWER real pairs than the
target, every real candidate found for that bin is kept and the script
moves on -- nothing errors, nothing is padded with fake data. The end
report shows target vs. actual for every one of the 50 bins.

HOW CANDIDATES ARE FOUND (mirrors the rules already verified elsewhere
in this project)
-----------------------------------------------------------------------
For every pair of DIFFERENT speakers of the SAME gender (A, B):
  - every one of A's enroll files crossed with every one of B's test
    files is a candidate, EXCEPT any specific pair that shares a
    meeting_id (that pair would violate the meeting-diff rule)
  - most speaker pairs never share a meeting at all (confirmed earlier:
    only ~8 gender-matched pairs out of ~2,367 do), so for the common
    case every combination is valid with no per-file check needed --
    only the rare overlapping pairs pay the extra check

Any candidate whose duration_diff would fall ABOVE the largest positive
bin edge is discarded (there's no positive-side bin to match it
against).

MEMORY, NOT JUST DISK
----------------------
The full negative candidate space is about 1.2 billion pairs -- far too
many to hold in memory or on disk. This script uses reservoir sampling
(the same idea discussed earlier): for each of the 50 bins, it keeps a
fixed-size holding area sized to that bin's target, and as candidates
stream past, each one has an equal chance of ending up in the final
sample, regardless of when in the run it was seen. Every candidate
pair IS computed (time is not the constraint here), but only the
capped, sampled subset is ever stored.

Internally, while streaming, each candidate is kept as three small
integers (an id for the enroll path, an id for the test path, and the
duration difference in centiseconds) -- using the exact same path
numbering as path_lookup.csv from the positive-trials encoding, so a
given file has the same id in both positive and negative data. This
keeps the running memory use to a few GB even at the full 273-million
upper bound. Only at the very end are those ids turned back into real
path text, when the final (much smaller, capped) table is built.

THIS WILL TAKE A WHILE
------------------------
Visiting ~1.2 billion candidate pairs in Python is not instant. Expect
somewhere between 15 minutes and a couple of hours depending on your
machine -- progress is printed periodically with an ETA based on the
already-known real total. Time was explicitly said not to be the
constraint here; this script is written to make one honest pass and
report exactly what it found, not to be clever about cutting the pass
short.

Usage:
    pip install pandas joblib --break-system-packages   (or via conda,
                                                           matching how
                                                           pandas itself
                                                           is installed)

    python create_negative_trials.py \\
        eval_xlsx_out/eval_enroll_test2.xlsx \\
        positive_trial_full_distribution_stats.json \\
        encoded_trials/path_lookup.csv \\
        --ratio 10 \\
        --out negative_trials.joblib
"""

import argparse
import bisect
import csv
import json
import random
import re
import sys
import time
from array import array
from pathlib import Path

DUR_RE = re.compile(r"_(?P<begin>\d+)-(?P<end>\d+)\.wav$", re.IGNORECASE)
REQUIRED_COLS = {"wav_path", "speaker_id", "meeting_id", "mic_type"}

# from the real exact-count run earlier in this project (Exact_Trial_Counts.json)
# -- used ONLY to print a progress ETA, has no effect on correctness
EXPECTED_TOTAL_CANDIDATES = 1_224_707_432


def shorten_path(wav_path: str, mic_type: str) -> str:
    marker = f"/{mic_type}/"
    idx = wav_path.find(marker)
    return wav_path[idx:] if idx != -1 else wav_path


def duration_cs_from_path(wav_path: str):
    """Same rule used everywhere else in this project:
    duration_s = (end - begin) / 100.0. Returned here as an integer
    number of CENTISECONDS (i.e. end - begin) so all downstream math
    stays in exact integers until the very last step."""
    m = DUR_RE.search(wav_path)
    if not m:
        return None
    return int(m.group("end")) - int(m.group("begin"))


def load_sheet(wb, sheet_name, path_to_id):
    """Returns dict: speaker_id -> list of (path_id, meeting_id, duration_cs)."""
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: no '{sheet_name}' sheet found. Sheets present: {wb.sheetnames}",
              file=sys.stderr)
        sys.exit(1)
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows_iter)]
    missing = REQUIRED_COLS - set(header)
    if missing:
        print(f"ERROR: sheet '{sheet_name}' missing columns {missing}.", file=sys.stderr)
        sys.exit(1)
    idx = {name: header.index(name) for name in header}

    by_speaker = {}
    n_no_duration = 0
    n_unknown_path = 0
    for row in rows_iter:
        if row is None:
            continue
        wav_path = row[idx["wav_path"]]
        sid = row[idx["speaker_id"]]
        meeting = row[idx["meeting_id"]]
        mic = row[idx["mic_type"]]
        if wav_path is None or sid is None or meeting is None:
            continue
        wav_path = str(wav_path)
        dur_cs = duration_cs_from_path(wav_path)
        if dur_cs is None:
            n_no_duration += 1
            continue
        short = shorten_path(wav_path, str(mic).strip().lower())
        path_id = path_to_id.get(short)
        if path_id is None:
            n_unknown_path += 1
            continue
        by_speaker.setdefault(str(sid), []).append((path_id, str(meeting), dur_cs))

    if n_no_duration:
        print(f"WARNING: {n_no_duration} rows in '{sheet_name}' had no parseable "
              f"begin-end timestamp -- skipped.", file=sys.stderr)
    if n_unknown_path:
        print(f"ERROR: {n_unknown_path} rows in '{sheet_name}' have a path not found "
              f"in path_lookup.csv. path_lookup.csv must come from encoding the SAME "
              f"positive_trials.csv that was built from this same xlsx -- "
              f"regenerate path_lookup.csv if the xlsx changed.", file=sys.stderr)
        sys.exit(1)
    return by_speaker


def load_path_lookup(lookup_path: Path):
    """Returns (path_to_id dict, id_to_path list) -- reusing the EXACT same
    numbering the positive trials already use, so a file has one id across
    both positive and negative data."""
    id_to_path = []
    path_to_id = {}
    with open(lookup_path, newline="") as f:
        reader = csv.reader(f)
        next(reader)
        for row in reader:
            if not row:
                continue
            pid, path = int(row[0]), row[1]
            while len(id_to_path) <= pid:
                id_to_path.append(None)
            id_to_path[pid] = path
            path_to_id[path] = pid
    return path_to_id, id_to_path


def load_bins(stats_json_path: Path, ratio: int):
    with open(stats_json_path) as f:
        stats = json.load(f)
    bins = stats["equal_width_bins"]["bins"]
    lowers = [b["lower_seconds"] for b in bins]
    uppers = [b["upper_seconds"] for b in bins]
    pos_counts = [b["count"] for b in bins]
    targets = [c * ratio for c in pos_counts]
    return lowers, uppers, pos_counts, targets


def find_bin(diff_s, lowers, max_upper):
    """Equal-width bins, half-open [lower, upper) except the very last bin
    which also includes its own upper edge (matches numpy.histogram's own
    convention, which is what produced these edges in the first place).
    Returns None if diff_s is out of range entirely (discard)."""
    if diff_s > max_upper:
        return None
    idx = bisect.bisect_right(lowers, diff_s) - 1
    if idx < 0:
        idx = 0
    return idx


class Reservoir:
    """One bin's fixed-capacity, unbiased random sample (Algorithm R),
    stored as three parallel int arrays -- not Python tuples/objects --
    to keep memory to ~12 bytes/row even at hundreds of millions of rows."""

    __slots__ = ("capacity", "seen", "enroll_ids", "test_ids", "diffs_cs")

    def __init__(self, capacity):
        self.capacity = capacity
        self.seen = 0
        self.enroll_ids = array("i")
        self.test_ids = array("i")
        self.diffs_cs = array("i")

    def offer(self, e_id, t_id, diff_cs):
        self.seen += 1
        if len(self.enroll_ids) < self.capacity:
            self.enroll_ids.append(e_id)
            self.test_ids.append(t_id)
            self.diffs_cs.append(diff_cs)
        else:
            j = random.randint(0, self.seen - 1)
            if j < self.capacity:
                self.enroll_ids[j] = e_id
                self.test_ids[j] = t_id
                self.diffs_cs[j] = diff_cs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="Path to eval_enroll_test2.xlsx (meeting-split workbook)")
    ap.add_argument("positive_stats_json",
                     help="Path to positive_trial_full_distribution_stats.json")
    ap.add_argument("path_lookup_csv",
                     help="Path to path_lookup.csv from encoding the positive trials")
    ap.add_argument("--ratio", type=int, default=10,
                     help="Target negatives per positive, per bin (default 10)")
    ap.add_argument("--out", default="negative_trials.joblib")
    ap.add_argument("--report", default="negative_trials_report.json")
    ap.add_argument("--progress-every", type=int, default=20_000_000,
                     help="Print progress every N candidate pairs (default 20,000,000)")
    ap.add_argument("--seed", type=int, default=42,
                     help="Random seed for reservoir sampling (default 42, for repeatability)")
    args = ap.parse_args()

    random.seed(args.seed)

    xlsx_path = Path(args.xlsx)
    stats_path = Path(args.positive_stats_json)
    lookup_path = Path(args.path_lookup_csv)
    for p in (xlsx_path, stats_path, lookup_path):
        if not p.is_file():
            print(f"ERROR: {p} not found.", file=sys.stderr)
            sys.exit(1)

    try:
        import pandas as pd
        import numpy as np
        import joblib
    except ImportError as e:
        print(f"Missing dependency ({e}). Install pandas, numpy, joblib the same way "
              f"you already confirmed working (conda, on this machine).", file=sys.stderr)
        sys.exit(1)

    print(f"Loading bin edges and targets from {stats_path}...", file=sys.stderr)
    lowers, uppers, pos_counts, targets = load_bins(stats_path, args.ratio)
    n_bins = len(lowers)
    max_upper = uppers[-1]
    print(f"  {n_bins} bins, 0.0s to {max_upper:.2f}s, ratio={args.ratio}x", file=sys.stderr)
    print(f"  total positive trials across all bins : {sum(pos_counts):,}", file=sys.stderr)
    print(f"  upper-bound target (sum of all bins)   : {sum(targets):,}", file=sys.stderr)

    print(f"\nLoading {lookup_path}...", file=sys.stderr)
    path_to_id, id_to_path = load_path_lookup(lookup_path)
    print(f"  {len(id_to_path):,} known paths", file=sys.stderr)

    from openpyxl import load_workbook
    print(f"\nLoading {xlsx_path}...", file=sys.stderr)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    enroll_by_speaker = load_sheet(wb, "Enroll", path_to_id)
    test_by_speaker = load_sheet(wb, "Test", path_to_id)
    wb.close()

    eligible = sorted(set(enroll_by_speaker) & set(test_by_speaker))
    print(f"  {len(eligible):,} eligible speakers (have both enroll and test rows)",
          file=sys.stderr)

    male = [s for s in eligible if s[0] == "m"]
    female = [s for s in eligible if s[0] == "f"]
    print(f"  {len(male)} male, {len(female)} female", file=sys.stderr)

    reservoirs = [Reservoir(t) for t in targets]

    n_candidates = 0
    n_discarded_out_of_range = 0
    n_pairs_fast_path = 0
    n_pairs_slow_path = 0
    t0 = time.time()

    def process_gender_group(speakers, label):
        nonlocal n_candidates, n_discarded_out_of_range
        nonlocal n_pairs_fast_path, n_pairs_slow_path

        # precompute each speaker's enroll-meeting-set and test-meeting-set once
        enroll_meetings = {s: {m for (_, m, _) in enroll_by_speaker[s]} for s in speakers}
        test_meetings = {s: {m for (_, m, _) in test_by_speaker[s]} for s in speakers}

        for a in speakers:
            e_files = enroll_by_speaker[a]
            for b in speakers:
                if a == b:
                    continue
                overlap = enroll_meetings[a] & test_meetings[b]
                t_files = test_by_speaker[b]

                if not overlap:
                    n_pairs_fast_path += 1
                    for (e_id, _e_m, e_dur) in e_files:
                        for (t_id, _t_m, t_dur) in t_files:
                            diff_cs = e_dur - t_dur
                            if diff_cs < 0:
                                diff_cs = -diff_cs
                            n_candidates += 1
                            b_idx = find_bin(diff_cs / 100.0, lowers, max_upper)
                            if b_idx is None:
                                n_discarded_out_of_range += 1
                            else:
                                reservoirs[b_idx].offer(e_id, t_id, diff_cs)
                            if n_candidates % args.progress_every == 0:
                                _print_progress(n_candidates, t0)
                else:
                    n_pairs_slow_path += 1
                    for (e_id, e_m, e_dur) in e_files:
                        for (t_id, t_m, t_dur) in t_files:
                            if e_m == t_m:
                                continue
                            diff_cs = e_dur - t_dur
                            if diff_cs < 0:
                                diff_cs = -diff_cs
                            n_candidates += 1
                            b_idx = find_bin(diff_cs / 100.0, lowers, max_upper)
                            if b_idx is None:
                                n_discarded_out_of_range += 1
                            else:
                                reservoirs[b_idx].offer(e_id, t_id, diff_cs)
                            if n_candidates % args.progress_every == 0:
                                _print_progress(n_candidates, t0)

        print(f"  {label}: done ({len(speakers)} speakers, "
              f"{n_pairs_fast_path + n_pairs_slow_path} ordered pairs so far)",
              file=sys.stderr)

    def _print_progress(n_done, t0):
        elapsed = time.time() - t0
        rate = n_done / elapsed if elapsed > 0 else 0
        pct = 100.0 * n_done / EXPECTED_TOTAL_CANDIDATES
        remaining = (EXPECTED_TOTAL_CANDIDATES - n_done) / rate if rate > 0 else float("inf")
        print(f"  ... {n_done:,} candidates processed "
              f"(~{pct:.1f}% of expected {EXPECTED_TOTAL_CANDIDATES:,}), "
              f"{rate:,.0f}/s, ~{remaining/60:.1f} min remaining", file=sys.stderr)

    print(f"\nStreaming candidate negative pairs "
          f"(~{EXPECTED_TOTAL_CANDIDATES:,} expected total)...", file=sys.stderr)
    process_gender_group(female, "female")
    process_gender_group(male, "male")

    elapsed_total = time.time() - t0
    print(f"\nDone streaming. {n_candidates:,} candidates seen "
          f"({n_discarded_out_of_range:,} discarded as out of range) "
          f"in {elapsed_total/60:.1f} minutes.", file=sys.stderr)
    print(f"  speaker pairs needing no per-file check (no shared meeting): "
          f"{n_pairs_fast_path:,}", file=sys.stderr)
    print(f"  speaker pairs needing per-file exclusion (shared a meeting): "
          f"{n_pairs_slow_path:,}", file=sys.stderr)

    # ---- flatten all 50 reservoirs into one set of columns ----
    print("\nFlattening reservoirs into final columns...", file=sys.stderr)
    all_enroll = array("i")
    all_test = array("i")
    all_diff_cs = array("i")
    bin_report = []
    for i in range(n_bins):
        r = reservoirs[i]
        got = len(r.enroll_ids)
        all_enroll.extend(r.enroll_ids)
        all_test.extend(r.test_ids)
        all_diff_cs.extend(r.diffs_cs)
        bin_report.append({
            "bin_index": i,
            "lower_seconds": lowers[i],
            "upper_seconds": uppers[i],
            "positive_count": pos_counts[i],
            "target_negative_count": targets[i],
            "actual_negative_count": got,
            "shortfall": max(0, targets[i] - got),
            "candidates_seen_for_this_bin": r.seen,
        })

    total_rows = len(all_enroll)
    print(f"  total negative trials kept: {total_rows:,} "
          f"(upper bound was {sum(targets):,})", file=sys.stderr)

    # ---- build the readable pandas table ----
    print("\nBuilding the pandas DataFrame (category dtype for paths)...", file=sys.stderr)
    enroll_codes = np.frombuffer(all_enroll, dtype=np.int32)
    test_codes = np.frombuffer(all_test, dtype=np.int32)
    diff_seconds = np.frombuffer(all_diff_cs, dtype=np.int32).astype(np.float32) / 100.0

    df = pd.DataFrame({
        "enroll_path": pd.Categorical.from_codes(enroll_codes, categories=id_to_path),
        "test_path": pd.Categorical.from_codes(test_codes, categories=id_to_path),
        "duration_diff": diff_seconds,
    })

    out_path = Path(args.out)
    print(f"Saving to {out_path} with joblib (compressed)...", file=sys.stderr)
    joblib.dump(df, out_path, compress=3)
    size_mb = out_path.stat().st_size / (1024 * 1024)
    print(f"  saved: {size_mb:.1f} MB, {len(df):,} rows", file=sys.stderr)

    report = {
        "xlsx": str(xlsx_path.resolve()),
        "ratio": args.ratio,
        "n_bins": n_bins,
        "candidates_seen": n_candidates,
        "candidates_discarded_out_of_range": n_discarded_out_of_range,
        "elapsed_minutes": round(elapsed_total / 60, 2),
        "total_negative_rows": total_rows,
        "upper_bound_target": sum(targets),
        "shortfall_total": sum(b["shortfall"] for b in bin_report),
        "output_file": str(out_path.resolve()),
        "output_size_mb": round(size_mb, 2),
        "bins": bin_report,
    }
    with open(args.report, "w") as f:
        json.dump(report, f, indent=2)
    print(f"Report written to {args.report}", file=sys.stderr)

    print("\nFirst 5 rows (exactly what df.head() shows after joblib.load):",
          file=sys.stderr)
    print(df.head().to_string(index=False), file=sys.stderr)


if __name__ == "__main__":
    main()