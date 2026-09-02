#!/usr/bin/env python3
#!/usr/bin/env python3
"""
ami_way_b_feasibility_check.py — checks whether Way B ("split each
speaker's meetings into two piles, one for enroll, one for test, so a
same-meeting collision is impossible") actually works on your real data,
before we build anything. Read-only: reads eval_enroll_test.xlsx (the
Enroll and Test sheets), writes nothing back, touches no audio files.

WHAT IT CHECKS, per eval speaker:
  - How many distinct meetings that speaker has.
  - Whether there's ANY way to split those meetings into an "enroll
    pile" and a "test pile" (no meeting in both) such that:
      - the enroll pile has >= 10 ihm files total
      - the test pile has >= 10 ihm+sdm files total
  - If yes: feasible, and it records one such split.
  - If no: infeasible, and it records WHY (too few meetings? all their
    files clustered in one meeting? not enough total files even
    ignoring the split?).

This directly tests the "con" we flagged for Way B: that a speaker might
come up short of 10 on one side once you force enroll/test into
non-overlapping meetings. Better to know the real number now than guess.

HOW THE CHECK WORKS (for the curious): for each speaker, every one of
their meetings could go to the "enroll pile" or the "test pile" — that's
2 choices per meeting, so trying every possible split is 2^(number of
meetings). AMI speakers only have a handful of meetings each, so this
brute-force is fast and exact (no approximation).

Output:
  way_b_feasibility.json   full per-speaker detail
  Printed summary: how many of the 100 eval speakers are feasible,
  how many aren't, and why.

Usage:
    python ami_way_b_feasibility_check.py eval_xlsx_out/eval_enroll_test.xlsx
"""

import argparse
import itertools
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_COLS = {"speaker_id", "meeting_id"}


def load_meeting_counts(xlsx_path: Path, sheet_name: str, wb):
    """Read one sheet via openpyxl only (no pandas anywhere) and return a
    dict: speaker_id -> Counter(meeting_id -> row count).

    This used to go through pandas (first pd.read_excel, then a plain
    pd.DataFrame() constructor) and both raised the identical 'TypeError:
    Cannot convert numpy.ndarray to numpy.ndarray'. The second traceback
    proved it wasn't an Excel-parsing bug at all -- it broke on
    pandas.DataFrame(plain_python_list, columns=plain_python_list_of_str),
    which has nothing to do with openpyxl or this file. That signature is
    a known symptom of pandas and numpy being binary-incompatible in this
    environment (pandas compiled against a different numpy C-ABI than the
    numpy actually installed) -- it can fail on almost ANY pandas array/
    Index construction, not just this script. Fix: don't ask pandas to
    build anything. The only thing this script needs is 'how many rows
    does each (speaker, meeting) pair have', which plain dicts/Counters
    do just fine without touching pandas at all."""
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows_iter)]
    missing = REQUIRED_COLS - set(header)
    if missing:
        print(f"ERROR: {sheet_name} sheet missing columns {missing}", file=sys.stderr)
        sys.exit(1)
    sid_idx = header.index("speaker_id")
    meeting_idx = header.index("meeting_id")

    by_speaker = defaultdict(Counter)
    for row in rows_iter:
        sid = row[sid_idx]
        meeting = row[meeting_idx]
        if sid is None or meeting is None:
            continue
        by_speaker[sid][str(meeting)] += 1
    return by_speaker


def load_sheets(xlsx_path: Path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    enroll_by_speaker = load_meeting_counts(xlsx_path, "Enroll", wb)
    test_by_speaker = load_meeting_counts(xlsx_path, "Test", wb)
    wb.close()
    return enroll_by_speaker, test_by_speaker


def check_speaker(enroll_meeting_counts, test_meeting_counts, need_enroll=10, need_test=10):
    """enroll_meeting_counts / test_meeting_counts: dict meeting_id -> count.
    Try every way to split the union of meetings into (enroll_side, test_side)
    with no meeting in both. Return the best result found."""
    meetings = sorted(set(enroll_meeting_counts) | set(test_meeting_counts))
    n = len(meetings)

    best = None  # (feasible, enroll_total, test_total, enroll_side, test_side)

    # brute force every 2^n assignment (n is small for real speakers)
    for bits in itertools.product([0, 1], repeat=n):
        enroll_side = [m for m, b in zip(meetings, bits) if b == 0]
        test_side = [m for m, b in zip(meetings, bits) if b == 1]
        e_total = sum(enroll_meeting_counts.get(m, 0) for m in enroll_side)
        t_total = sum(test_meeting_counts.get(m, 0) for m in test_side)
        feasible = e_total >= need_enroll and t_total >= need_test
        score = min(e_total, need_enroll) + min(t_total, need_test)  # how close to feasible
        if best is None or (feasible and not best[0]) or (feasible == best[0] and score > best[5]):
            best = (feasible, e_total, t_total, enroll_side, test_side, score)
        if feasible:
            # good enough once we find ANY feasible split; keep it
            return {
                "n_meetings": n, "feasible": True,
                "enroll_side_meetings": enroll_side, "test_side_meetings": test_side,
                "enroll_side_total": e_total, "test_side_total": t_total,
            }

    feasible, e_total, t_total, enroll_side, test_side, _ = best
    total_enroll_all = sum(enroll_meeting_counts.values())
    total_test_all = sum(test_meeting_counts.values())
    reason = []
    if n < 2:
        reason.append("only 1 meeting total -- can't split at all")
    if total_enroll_all < need_enroll:
        reason.append(f"speaker has only {total_enroll_all} ihm files total, need {need_enroll}")
    if total_test_all < need_test:
        reason.append(f"speaker has only {total_test_all} ihm+sdm files total, need {need_test}")
    if not reason:
        reason.append("has enough files overall, but they're too concentrated in one meeting to split cleanly")

    return {
        "n_meetings": n, "feasible": False,
        "best_enroll_side_meetings": enroll_side, "best_test_side_meetings": test_side,
        "best_enroll_side_total": e_total, "best_test_side_total": t_total,
        "total_enroll_all_meetings": total_enroll_all, "total_test_all_meetings": total_test_all,
        "reason": "; ".join(reason),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="Path to eval_enroll_test.xlsx")
    ap.add_argument("--need-enroll", type=int, default=10)
    ap.add_argument("--need-test", type=int, default=10)
    ap.add_argument("--out", default="way_b_feasibility.json")
    args = ap.parse_args()

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.is_file():
        print(f"ERROR: {xlsx_path} not found.", file=sys.stderr)
        sys.exit(1)

    enroll_by_speaker, test_by_speaker = load_sheets(xlsx_path)

    all_speakers = sorted(set(enroll_by_speaker) | set(test_by_speaker))

    results = {}
    for sid in all_speakers:
        results[sid] = check_speaker(
            enroll_by_speaker.get(sid, {}), test_by_speaker.get(sid, {}),
            need_enroll=args.need_enroll, need_test=args.need_test,
        )

    n_feasible = sum(1 for r in results.values() if r["feasible"])
    n_infeasible = len(results) - n_feasible
    infeasible_speakers = [sid for sid, r in results.items() if not r["feasible"]]
    reasons = Counter(results[sid]["reason"] for sid in infeasible_speakers)

    summary = {
        "input_xlsx": str(xlsx_path.resolve()),
        "need_enroll": args.need_enroll,
        "need_test": args.need_test,
        "total_speakers_checked": len(results),
        "feasible_count": n_feasible,
        "infeasible_count": n_infeasible,
        "infeasible_speakers": infeasible_speakers,
        "infeasible_reasons_breakdown": dict(reasons),
    }

    with open(args.out, "w") as f:
        json.dump({"summary": summary, "per_speaker": results}, f, indent=2)

    print(json.dumps(summary, indent=2))
    print(f"\nFull per-speaker detail written to {args.out}")


if __name__ == "__main__":
    main()