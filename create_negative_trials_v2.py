#!/usr/bin/env python3
"""
create_negative_trials_v2.py

Same job as create_negative_trials.py (capped NEGATIVE trial table, 50
duration_diff bins, ratio x positives per bin, reservoir sampling so the
~1.2 billion candidate pool is never fully stored) -- with one addition:
inside EACH of the 50 bins, negatives are now split into ihm-ihm and
ihm-sdm using the SAME ratio those two mic-combos have in the real
positive trials for that bin. The old script picked whichever candidates
happened to land in the reservoir first; this one guarantees the split.

WHY A SEPARATE FILE
--------------------
This is kept as v2 rather than overwriting create_negative_trials.py so
the original run (already completed, report already analysed) is not
touched. Nothing about the meeting-diff rule, the gender split, or the
overall bin edges/targets changes -- only how each bin's target capacity
is divided between the two mic-combos.

REQUIRES A NEW INPUT: positive_mic_ratio_stats.json
-----------------------------------------------------
Produced by compute_positive_mic_ratio.py, run once against
positive_trials.csv. It records, per bin, how many real positive trials
were ihm-ihm vs ihm-sdm. Run that script first if you haven't already --
this script will refuse to start without it.

HOW THE PER-BIN SPLIT WORKS
------------------------------
For bin i:
    total_target[i]     = positive_count[i] * ratio      (unchanged from v1)
    target_ihm_ihm[i]   = ihm_ihm_count[i]  * ratio
    target_ihm_sdm[i]   = ihm_sdm_count[i]  * ratio
    target_ihm_ihm[i] + target_ihm_sdm[i]  ==  total_target[i]   (always)

Each bin gets TWO reservoirs instead of one (100 total instead of 50),
each sized to its own sub-target. A candidate is routed to the ihm-ihm
reservoir or the ihm-sdm reservoir for its bin based on the TEST file's
mic type (enroll is always ihm, so enroll never decides this).

Everything else -- the meeting-diff fast/slow path, the reservoir
sampling algorithm itself, the final pandas category-dtype table, the
joblib save -- is identical to v1.

Usage:
    python create_negative_trials_v2.py \\
        eval_xlsx_out/eval_enroll_test2.xlsx \\
        positive_trial_full_distribution_stats.json \\
        positive_mic_ratio_stats.json \\
        encoded_trials/path_lookup.csv \\
        --ratio 10 \\
        --out negative_trials_v2.joblib
"""

import argparse
import bisect
import csv
import json
import random
import re
import sys
import time
from array import array
from pathlib import Path

DUR_RE = re.compile(r"_(?P<begin>\d+)-(?P<end>\d+)\.wav$", re.IGNORECASE)
REQUIRED_COLS = {"wav_path", "speaker_id", "meeting_id", "mic_type"}

# from the real exact-count run earlier in this project -- used ONLY for
# the printed progress ETA, has no effect on correctness
EXPECTED_TOTAL_CANDIDATES = 1_224_707_432


def shorten_path(wav_path: str, mic_type: str) -> str:
    marker = f"/{mic_type}/"
    idx = wav_path.find(marker)
    return wav_path[idx:] if idx != -1 else wav_path


def duration_cs_from_path(wav_path: str):
    m = DUR_RE.search(wav_path)
    if not m:
        return None
    return int(m.group("end")) - int(m.group("begin"))


def load_sheet(wb, sheet_name, path_to_id):
    """Returns dict: speaker_id -> list of (path_id, meeting_id, duration_cs, mic).
    'mic' is kept now (v1 dropped it) so the negative-generation loop can
    tell an ihm test file from an sdm test file."""
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: no '{sheet_name}' sheet found. Sheets present: {wb.sheetnames}",
              file=sys.stderr)
        sys.exit(1)
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows_iter)]
    missing = REQUIRED_COLS - set(header)
    if missing:
        print(f"ERROR: sheet '{sheet_name}' missing columns {missing}.", file=sys.stderr)
        sys.exit(1)
    idx = {name: header.index(name) for name in header}

    by_speaker = {}
    n_no_duration = 0
    n_unknown_path = 0
    for row in rows_iter:
        if row is None:
            continue
        wav_path = row[idx["wav_path"]]
        sid = row[idx["speaker_id"]]
        meeting = row[idx["meeting_id"]]
        mic = row[idx["mic_type"]]
        if wav_path is None or sid is None or meeting is None:
            continue
        wav_path = str(wav_path)
        mic_norm = str(mic).strip().lower()
        dur_cs = duration_cs_from_path(wav_path)
        if dur_cs is None:
            n_no_duration += 1
            continue
        short = shorten_path(wav_path, mic_norm)
        path_id = path_to_id.get(short)
        if path_id is None:
            n_unknown_path += 1
            continue
        by_speaker.setdefault(str(sid), []).append(
            (path_id, str(meeting), dur_cs, mic_norm))

    if n_no_duration:
        print(f"WARNING: {n_no_duration} rows in '{sheet_name}' had no parseable "
              f"begin-end timestamp -- skipped.", file=sys.stderr)
    if n_unknown_path:
        print(f"ERROR: {n_unknown_path} rows in '{sheet_name}' have a path not found "
              f"in path_lookup.csv. path_lookup.csv must come from encoding the SAME "
              f"positive_trials.csv that was built from this same xlsx -- "
              f"regenerate path_lookup.csv if the xlsx changed.", file=sys.stderr)
        sys.exit(1)
    return by_speaker


def load_path_lookup(lookup_path: Path):
    id_to_path = []
    path_to_id = {}
    with open(lookup_path, newline="") as f:
        reader = csv.reader(f)
        next(reader)
        for row in reader:
            if not row:
                continue
            pid, path = int(row[0]), row[1]
            while len(id_to_path) <= pid:
                id_to_path.append(None)
            id_to_path[pid] = path
            path_to_id[path] = pid
    return path_to_id, id_to_path


def load_bins_and_mic_targets(stats_json_path: Path, mic_ratio_json_path: Path, ratio: int):
    with open(stats_json_path) as f:
        stats = json.load(f)
    bins = stats["equal_width_bins"]["bins"]
    lowers = [b["lower_seconds"] for b in bins]
    uppers = [b["upper_seconds"] for b in bins]
    pos_counts = [b["count"] for b in bins]

    with open(mic_ratio_json_path) as f:
        mic_stats = json.load(f)
    mic_bins = mic_stats["bins"]

    n_bins = len(lowers)
    if len(mic_bins) != n_bins:
        print(f"ERROR: bin count mismatch -- {stats_json_path} has {n_bins} bins, "
              f"{mic_ratio_json_path} has {len(mic_bins)}. Regenerate "
              f"positive_mic_ratio_stats.json against the current stats file.",
              file=sys.stderr)
        sys.exit(1)

    pos_ihm_ihm = [0] * n_bins
    pos_ihm_sdm = [0] * n_bins
    for mb in mic_bins:
        i = mb["bin_index"]
        if abs(mb["lower_seconds"] - lowers[i]) > 1e-6 or abs(mb["upper_seconds"] - uppers[i]) > 1e-6:
            print(f"ERROR: bin edges don't line up at index {i} between the two "
                  f"stats files. Regenerate positive_mic_ratio_stats.json against "
                  f"the current positive_trial_full_distribution_stats.json.",
                  file=sys.stderr)
            sys.exit(1)
        pos_ihm_ihm[i] = mb["ihm_ihm_count"]
        pos_ihm_sdm[i] = mb["ihm_sdm_count"]

    target_ihm_ihm = [c * ratio for c in pos_ihm_ihm]
    target_ihm_sdm = [c * ratio for c in pos_ihm_sdm]
    return lowers, uppers, pos_counts, pos_ihm_ihm, pos_ihm_sdm, target_ihm_ihm, target_ihm_sdm


def find_bin(diff_s, lowers, max_upper):
    if diff_s > max_upper:
        return None
    idx = bisect.bisect_right(lowers, diff_s) - 1
    if idx < 0:
        idx = 0
    return idx


class Reservoir:
    """One (bin, mic-combo) pair's fixed-capacity, unbiased random sample
    (Algorithm R), stored as three parallel int arrays -- not Python
    tuples/objects -- to keep memory small even at hundreds of millions
    of rows total across all reservoirs."""

    __slots__ = ("capacity", "seen", "enroll_ids", "test_ids", "diffs_cs")

    def __init__(self, capacity):
        self.capacity = capacity
        self.seen = 0
        self.enroll_ids = array("i")
        self.test_ids = array("i")
        self.diffs_cs = array("i")

    def offer(self, e_id, t_id, diff_cs):
        self.seen += 1
        if len(self.enroll_ids) < self.capacity:
            self.enroll_ids.append(e_id)
            self.test_ids.append(t_id)
            self.diffs_cs.append(diff_cs)
        else:
            j = random.randint(0, self.seen - 1)
            if j < self.capacity:
                self.enroll_ids[j] = e_id
                self.test_ids[j] = t_id
                self.diffs_cs[j] = diff_cs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="Path to eval_enroll_test2.xlsx (meeting-split workbook)")
    ap.add_argument("positive_stats_json",
                     help="Path to positive_trial_full_distribution_stats.json")
    ap.add_argument("positive_mic_ratio_json",
                     help="Path to positive_mic_ratio_stats.json "
                          "(from compute_positive_mic_ratio.py)")
    ap.add_argument("path_lookup_csv",
                     help="Path to path_lookup.csv from encoding the positive trials")
    ap.add_argument("--ratio", type=int, default=10,
                     help="Target negatives per positive, per bin (default 10)")
    ap.add_argument("--out", default="negative_trials_v2.joblib")
    ap.add_argument("--report", default="negative_trials_v2_report.json")
    ap.add_argument("--progress-every", type=int, default=20_000_000)
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()

    random.seed(args.seed)

    xlsx_path = Path(args.xlsx)
    stats_path = Path(args.positive_stats_json)
    mic_ratio_path = Path(args.positive_mic_ratio_json)
    lookup_path = Path(args.path_lookup_csv)
    for p in (xlsx_path, stats_path, mic_ratio_path, lookup_path):
        if not p.is_file():
            print(f"ERROR: {p} not found.", file=sys.stderr)
            sys.exit(1)

    try:
        import pandas as pd
        import numpy as np
        import joblib
    except ImportError as e:
        print(f"Missing dependency ({e}). Install pandas, numpy, joblib the same way "
              f"you already confirmed working (conda, on this machine).", file=sys.stderr)
        sys.exit(1)

    print(f"Loading bin edges, positive counts, and mic-combo split...", file=sys.stderr)
    (lowers, uppers, pos_counts,
     pos_ihm_ihm, pos_ihm_sdm,
     target_ihm_ihm, target_ihm_sdm) = load_bins_and_mic_targets(
        stats_path, mic_ratio_path, args.ratio)
    n_bins = len(lowers)
    max_upper = uppers[-1]
    print(f"  {n_bins} bins, 0.0s to {max_upper:.2f}s, ratio={args.ratio}x", file=sys.stderr)
    print(f"  total positive trials             : {sum(pos_counts):,}", file=sys.stderr)
    print(f"  total upper-bound target (both combos combined) : "
          f"{sum(target_ihm_ihm) + sum(target_ihm_sdm):,}", file=sys.stderr)

    print(f"\nLoading {lookup_path}...", file=sys.stderr)
    path_to_id, id_to_path = load_path_lookup(lookup_path)
    print(f"  {len(id_to_path):,} known paths", file=sys.stderr)

    from openpyxl import load_workbook
    print(f"\nLoading {xlsx_path}...", file=sys.stderr)
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    enroll_by_speaker = load_sheet(wb, "Enroll", path_to_id)
    test_by_speaker = load_sheet(wb, "Test", path_to_id)
    wb.close()

    eligible = sorted(set(enroll_by_speaker) & set(test_by_speaker))
    print(f"  {len(eligible):,} eligible speakers (have both enroll and test rows)",
          file=sys.stderr)

    male = [s for s in eligible if s[0] == "m"]
    female = [s for s in eligible if s[0] == "f"]
    print(f"  {len(male)} male, {len(female)} female", file=sys.stderr)

    reservoirs_ihm = [Reservoir(t) for t in target_ihm_ihm]
    reservoirs_sdm = [Reservoir(t) for t in target_ihm_sdm]

    n_candidates = 0
    n_discarded_out_of_range = 0
    n_pairs_fast_path = 0
    n_pairs_slow_path = 0
    t0 = time.time()

    def route(e_id, t_id, diff_cs, t_mic):
        nonlocal n_discarded_out_of_range
        b_idx = find_bin(diff_cs / 100.0, lowers, max_upper)
        if b_idx is None:
            n_discarded_out_of_range += 1
            return
        if t_mic == "ihm":
            reservoirs_ihm[b_idx].offer(e_id, t_id, diff_cs)
        else:
            reservoirs_sdm[b_idx].offer(e_id, t_id, diff_cs)

    def process_gender_group(speakers, label):
        nonlocal n_candidates, n_pairs_fast_path, n_pairs_slow_path

        enroll_meetings = {s: {m for (_, m, _, _) in enroll_by_speaker[s]} for s in speakers}
        test_meetings = {s: {m for (_, m, _, _) in test_by_speaker[s]} for s in speakers}

        for a in speakers:
            e_files = enroll_by_speaker[a]
            for b in speakers:
                if a == b:
                    continue
                overlap = enroll_meetings[a] & test_meetings[b]
                t_files = test_by_speaker[b]

                if not overlap:
                    n_pairs_fast_path += 1
                    for (e_id, _e_m, e_dur, _e_mic) in e_files:
                        for (t_id, _t_m, t_dur, t_mic) in t_files:
                            diff_cs = e_dur - t_dur
                            if diff_cs < 0:
                                diff_cs = -diff_cs
                            n_candidates += 1
                            route(e_id, t_id, diff_cs, t_mic)
                            if n_candidates % args.progress_every == 0:
                                _print_progress(n_candidates, t0)
                else:
                    n_pairs_slow_path += 1
                    for (e_id, e_m, e_dur, _e_mic) in e_files:
                        for (t_id, t_m, t_dur, t_mic) in t_files:
                            if e_m == t_m:
                                continue
                            diff_cs = e_dur - t_dur
                            if diff_cs < 0:
                                diff_cs = -diff_cs
                            n_candidates += 1
                            route(e_id, t_id, diff_cs, t_mic)
                            if n_candidates % args.progress_every == 0:
                                _print_progress(n_candidates, t0)

        print(f"  {label}: done ({len(speakers)} speakers, "
              f"{n_pairs_fast_path + n_pairs_slow_path} ordered pairs so far)",
              file=sys.stderr)

    def _print_progress(n_done, t0):
        elapsed = time.time() - t0
        rate = n_done / elapsed if elapsed > 0 else 0
        pct = 100.0 * n_done / EXPECTED_TOTAL_CANDIDATES
        remaining = (EXPECTED_TOTAL_CANDIDATES - n_done) / rate if rate > 0 else float("inf")
        print(f"  ... {n_done:,} candidates processed "
              f"(~{pct:.1f}% of expected {EXPECTED_TOTAL_CANDIDATES:,}), "
              f"{rate:,.0f}/s, ~{remaining/60:.1f} min remaining", file=sys.stderr)

    print(f"\nStreaming candidate negative pairs "
          f"(~{EXPECTED_TOTAL_CANDIDATES:,} expected total)...", file=sys.stderr)
    process_gender_group(female, "female")
    process_gender_group(male, "male")

    elapsed_total = time.time() - t0
    print(f"\nDone streaming. {n_candidates:,} candidates seen "
          f"({n_discarded_out_of_range:,} discarded as out of range) "
          f"in {elapsed_total/60:.1f} minutes.", file=sys.stderr)
    print(f"  speaker pairs needing no per-file check (no shared meeting): "
          f"{n_pairs_fast_path:,}", file=sys.stderr)
    print(f"  speaker pairs needing per-file exclusion (shared a meeting): "
          f"{n_pairs_slow_path:,}", file=sys.stderr)

    print("\nFlattening reservoirs into final columns...", file=sys.stderr)
    all_enroll = array("i")
    all_test = array("i")
    all_diff_cs = array("i")
    bin_report = []
    for i in range(n_bins):
        r_ihm = reservoirs_ihm[i]
        r_sdm = reservoirs_sdm[i]
        got_ihm = len(r_ihm.enroll_ids)
        got_sdm = len(r_sdm.enroll_ids)

        all_enroll.extend(r_ihm.enroll_ids)
        all_test.extend(r_ihm.test_ids)
        all_diff_cs.extend(r_ihm.diffs_cs)

        all_enroll.extend(r_sdm.enroll_ids)
        all_test.extend(r_sdm.test_ids)
        all_diff_cs.extend(r_sdm.diffs_cs)

        bin_report.append({
            "bin_index": i,
            "lower_seconds": lowers[i],
            "upper_seconds": uppers[i],
            "positive_total": pos_counts[i],
            "positive_ihm_ihm": pos_ihm_ihm[i],
            "positive_ihm_sdm": pos_ihm_sdm[i],
            "target_ihm_ihm": target_ihm_ihm[i],
            "target_ihm_sdm": target_ihm_sdm[i],
            "actual_ihm_ihm": got_ihm,
            "actual_ihm_sdm": got_sdm,
            "actual_total": got_ihm + got_sdm,
            "shortfall_ihm_ihm": max(0, target_ihm_ihm[i] - got_ihm),
            "shortfall_ihm_sdm": max(0, target_ihm_sdm[i] - got_sdm),
            "candidates_seen_ihm_ihm": r_ihm.seen,
            "candidates_seen_ihm_sdm": r_sdm.seen,
        })

    total_rows = len(all_enroll)
    upper_bound = sum(target_ihm_ihm) + sum(target_ihm_sdm)
    print(f"  total negative trials kept: {total_rows:,} "
          f"(upper bound was {upper_bound:,})", file=sys.stderr)

    print("\nBuilding the pandas DataFrame (category dtype for paths)...", file=sys.stderr)
    enroll_codes = np.frombuffer(all_enroll, dtype=np.int32)
    test_codes = np.frombuffer(all_test, dtype=np.int32)
    diff_seconds = np.frombuffer(all_diff_cs, dtype=np.int32).astype(np.float32) / 100.0

    df = pd.DataFrame({
        "enroll_path": pd.Categorical.from_codes(enroll_codes, categories=id_to_path),
        "test_path": pd.Categorical.from_codes(test_codes, categories=id_to_path),
        "duration_diff": diff_seconds,
    })

    out_path = Path(args.out)
    print(f"Saving to {out_path} with joblib (compressed)...", file=sys.stderr)
    joblib.dump(df, out_path, compress=3)
    size_mb = out_path.stat().st_size / (1024 * 1024)
    print(f"  saved: {size_mb:.1f} MB, {len(df):,} rows", file=sys.stderr)

    report = {
        "xlsx": str(xlsx_path.resolve()),
        "ratio": args.ratio,
        "n_bins": n_bins,
        "candidates_seen": n_candidates,
        "candidates_discarded_out_of_range": n_discarded_out_of_range,
        "elapsed_minutes": round(elapsed_total / 60, 2),
        "total_negative_rows": total_rows,
        "upper_bound_target": upper_bound,
        "shortfall_total": sum(b["shortfall_ihm_ihm"] + b["shortfall_ihm_sdm"] for b in bin_report),
        "output_file": str(out_path.resolve()),
        "output_size_mb": round(size_mb, 2),
        "bins": bin_report,
    }
    with open(args.report, "w") as f:
        json.dump(report, f, indent=2)
    print(f"Report written to {args.report}", file=sys.stderr)

    print("\nFirst 5 rows (exactly what df.head() shows after joblib.load):",
          file=sys.stderr)
    print(df.head().to_string(index=False), file=sys.stderr)


if __name__ == "__main__":
    main()