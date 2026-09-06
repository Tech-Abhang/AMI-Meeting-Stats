#!/usr/bin/env python3
"""
ami_eval_enroll_test_xlsx.py — builds the eval-group ENROLL/TEST Excel
sheet per your latest spec. Does NOT touch, move, or rename anything on
disk — reads the existing ihm/ and sdm/ folders only, writes one new
.xlsx with wav file paths + metadata for future mapping.

LOCKED-IN SPEC (v4 — meeting-split rework, replaces v3's "same 100
speakers in both Enroll and Test" design):

  1. Eval group = 100 of 190 speakers: 70 male + 30 female, drawn
     randomly (seed 42, reproducible) from the full pool of 130 male /
     60 female speakers. The other ~90 speakers go into a separate
     "train group" speaker list — NOT processed further, per your
     standing instruction to defer train. UNCHANGED from v3.

  2. THE CHANGE — per eval speaker, split THEIR MEETINGS (not their
     files) in half:
       - Sort that speaker's distinct meeting_ids alphabetically.
       - The first floor(n/2) meetings -> ENROLL side.
       - The remaining ceil(n/2) meetings -> TEST side.
     ENROLL = ihm files ONLY, from the enroll-side meetings.
     TEST   = ihm + sdm files, from the TEST-side meetings (the other
              half — never the same meetings as Enroll).

     Worked example from your message: fee051 has meeting A, meeting B.
       Enroll: ihm, meeting A
       Test:   ihm meeting B + sdm meeting B

     This means Enroll and Test meetings are NEVER the same for a given
     speaker, by construction — the meeting-diff rule from your trial-pair
     spec is now baked into the data itself, not something you have to
     filter for afterward.

  3. Odd number of meetings (e.g. 3, 5): confirmed by you — Enroll gets
     the SMALLER half (floor(n/2)), Test gets the larger half (ceil(n/2)).
     E.g. 3 meetings -> 1 to Enroll, 2 to Test. 5 meetings -> 2 to
     Enroll, 3 to Test.

  4. Speakers with only 1 meeting total: confirmed by you — SKIP THEM
     ENTIRELY. There's no way to give a 1-meeting speaker a separate
     enroll meeting and test meeting, so they get zero rows in both the
     Enroll and Test sheets. On the real corpus this was 9 of the 100
     eval speakers (fee085, fee096, feo084, mee075, mee076, meo074,
     mio066, mio098, mio105) — the other 91 go through the new split.
     These 9 are listed separately in the Summary sheet and in the new
     "Dropped_Speakers" sheet so nothing silently disappears.

  This is a real behavior change from v3: v3 put a speaker's ihm files
  in Enroll AND (as part of ihm+sdm) also in Test — deliberate overlap,
  confirmed fine by you at the time. v4 does the opposite: Enroll and
  Test never share a meeting for the same speaker, and 9 speakers who
  can't support that split are excluded rather than force-included.

Output (.xlsx, written next to this script unless --out-dir is given):
  Sheet "Enroll"              — one row per ihm utterance, enroll-side
                                 meetings only, the 91 eligible eval speakers
  Sheet "Test"                — one row per ihm+sdm utterance, test-side
                                 meetings only, the same 91 speakers
  Sheet "Meeting_Split_Detail" — audit sheet: for each of the 91 speakers,
                                 which meetings went to which side, and how
                                 many rows that produced — so you can check
                                 any single speaker's split without recounting
                                 thousands of rows by hand
  Sheet "Dropped_Speakers"    — the 9 (on real data) skipped 1-meeting
                                 speakers, so they're documented, not lost
  Sheet "Train_Group_Speakers" — speaker-level list of the deferred ~90
                                 (unrelated to this change, untouched)
  Sheet "Summary"             — counts, some via COUNTIFS formulas so they
                                 stay correct if you manually prune rows
                                 later in Excel

Columns (Enroll/Test sheets):
  wav_path, speaker_id, gender, location, language, participant_number,
  role_suffix, native_split, meeting_id, mic_type

Usage:
    pip install openpyxl --break-system-packages
    python ami_eval_enroll_test_xlsx.py "/media/iiitd/My Passport Sachin/PhoneBus/data/AMI_Meeting_Corpus/audio" \
        --out-dir ./eval_xlsx_out
"""

import argparse
import json
import random
import re
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Missing dependency. Run: pip install openpyxl --break-system-packages", file=sys.stderr)
    raise

STEM_RE = re.compile(
    r"^(?P<sid>[a-zA-Z0-9]+)_(?P<meeting>[a-zA-Z0-9]+)_"
    r"(?P<split_in_name>train|dev|eval)_(?P<miccode>[a-zA-Z0-9]+)_"
    r"(?P<begin>\d+)-(?P<end>\d+)$"
)
NOSEG_RE = re.compile(
    r"^(?P<sid>[a-zA-Z0-9]+)_(?P<meeting>[a-zA-Z0-9]+)_"
    r"(?P<split_in_name>train|dev|eval)_(?P<miccode>[a-zA-Z0-9]+)$"
)

LOC = {"i": "Idiap", "e": "Edinburgh", "t": "TNO"}
LANG = {"e": "English", "d": "Dutch", "o": "Other"}
SID_TAIL_RE = re.compile(r"^(\d+)([a-zA-Z]*)$")

SEED = 42
EVAL_MALE_TARGET = 70
EVAL_FEMALE_TARGET = 30
FONT_NAME = "Arial"


def decode_speaker_id(sid: str):
    if len(sid) < 4:
        return None
    g, loc, lang = sid[0].lower(), sid[1].lower(), sid[2].lower()
    rest = sid[3:]
    if g not in ("f", "m") or loc not in LOC or lang not in LANG:
        return None
    m = SID_TAIL_RE.match(rest)
    if not m:
        return None
    return {
        "gender": "Female" if g == "f" else "Male",
        "location": LOC[loc],
        "language": LANG[lang],
        "number": m.group(1),
        "role_suffix": m.group(2) or None,
    }


def find_child_ci(parent: Path, name: str):
    if not parent.is_dir():
        return None
    target = name.lower()
    for child in parent.iterdir():
        if child.is_dir() and child.name.lower() == target:
            return child
    return None


def locate_audio_root(user_path: Path):
    candidates = [user_path, user_path / "audio", user_path.parent]
    tried = []
    for cand in candidates:
        if not cand.is_dir():
            tried.append((cand, "does not exist / not a directory"))
            continue
        ihm = find_child_ci(cand, "ihm")
        sdm = find_child_ci(cand, "sdm")
        top_level = sorted(p.name for p in cand.iterdir())[:25]
        tried.append((cand, f"ihm found: {bool(ihm)}, sdm found: {bool(sdm)}; top-level: {top_level}"))
        if ihm and sdm:
            return cand, ihm, sdm
    print("ERROR: could not locate both an 'ihm' and an 'sdm' folder from the path you gave.", file=sys.stderr)
    for cand, msg in tried:
        print(f"  - {cand}  ->  {msg}", file=sys.stderr)
    return None, None, None


def walk_mic_dir(mic_dir: Path, mic_label: str):
    rows = []
    for split_name in ("train", "dev", "eval"):
        split_dir = find_child_ci(mic_dir, split_name)
        if split_dir is None:
            continue
        for speaker_dir in sorted(p for p in split_dir.iterdir() if p.is_dir()):
            speaker_id = speaker_dir.name
            for wav_path in speaker_dir.iterdir():
                if not (wav_path.is_file() and wav_path.suffix.lower() == ".wav"):
                    continue
                stem = wav_path.stem
                m = STEM_RE.match(stem) or NOSEG_RE.match(stem)
                meeting_id = m.group("meeting") if m else None
                rows.append({
                    "wav_path": str(wav_path.resolve()),
                    "speaker_id": speaker_id,
                    "native_split": split_dir.name,
                    "mic_type": mic_label,
                    "meeting_id": meeting_id,
                    "parse_ok": bool(m),
                })
    return rows


def compute_meeting_splits(ihm_rows, eval_ids):
    """For each speaker in eval_ids, look at the distinct meeting_ids that
    show up in their ihm files (meeting_id=None rows -- parse failures --
    are ignored here, same as everywhere else in this script). Sort those
    meeting_ids and cut them in half:
        enroll side = first floor(n/2) meetings
        test side   = remaining ceil(n/2) meetings
    Speakers with fewer than 2 distinct meetings are skipped entirely
    (returned in `skipped`, not in `splits`) -- confirmed behavior.

    Returns:
        splits: speaker_id -> {"enroll_meetings": [...], "test_meetings": [...], "n_meetings": n}
        skipped: sorted list of speaker_ids with < 2 meetings
    """
    meetings_by_speaker = defaultdict(set)
    for r in ihm_rows:
        if r["speaker_id"] in eval_ids and r["meeting_id"] is not None:
            meetings_by_speaker[r["speaker_id"]].add(r["meeting_id"])

    splits = {}
    skipped = []
    for sid in sorted(eval_ids):
        meetings = sorted(meetings_by_speaker.get(sid, set()))
        n = len(meetings)
        if n < 2:
            skipped.append(sid)
            continue
        n_enroll = n // 2  # floor -> Enroll gets the smaller (or equal) half, confirmed
        splits[sid] = {
            "enroll_meetings": meetings[:n_enroll],
            "test_meetings": meetings[n_enroll:],
            "n_meetings": n,
        }
    return splits, skipped


def style_header(ws, ncols):
    header_font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
    ws.freeze_panes = "A2"


def autosize(ws, ncols, sample_rows=200):
    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        max_len = 0
        for r, cell in enumerate(ws[col_letter]):
            if r > sample_rows:
                break
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)


def write_data_sheet(wb, name, rows, columns):
    ws = wb.create_sheet(name)
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c, "") if r.get(c) is not None else "" for c in columns])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT_NAME)
    style_header(ws, len(columns))
    autosize(ws, len(columns))
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("root", help="Path to (or near) the 'audio' folder containing ihm/ and sdm/")
    ap.add_argument("--out-dir", default=".")
    args = ap.parse_args()

    resolved_root, ihm_dir, sdm_dir = locate_audio_root(Path(args.root))
    if resolved_root is None:
        sys.exit(1)
    print(f"Using root: {resolved_root}\n  ihm -> {ihm_dir}\n  sdm -> {sdm_dir}", file=sys.stderr)

    ihm_rows = walk_mic_dir(ihm_dir, "ihm")
    sdm_rows = walk_mic_dir(sdm_dir, "sdm")
    all_rows = ihm_rows + sdm_rows
    if not all_rows:
        print("ERROR: found ihm/sdm folders but zero .wav files under them. Nothing to do.", file=sys.stderr)
        sys.exit(1)

    all_speaker_ids = sorted(set(r["speaker_id"] for r in all_rows))
    decode_map = {}
    decode_failures = []
    for sid in all_speaker_ids:
        dec = decode_speaker_id(sid)
        if dec is None:
            decode_failures.append(sid)
        else:
            decode_map[sid] = dec

    male_ids = sorted(sid for sid, d in decode_map.items() if d["gender"] == "Male")
    female_ids = sorted(sid for sid, d in decode_map.items() if d["gender"] == "Female")

    if len(male_ids) < EVAL_MALE_TARGET or len(female_ids) < EVAL_FEMALE_TARGET:
        print(f"ERROR: not enough speakers to hit the target split. "
              f"Found {len(male_ids)} male (need {EVAL_MALE_TARGET}), "
              f"{len(female_ids)} female (need {EVAL_FEMALE_TARGET}).", file=sys.stderr)
        sys.exit(1)

    rng = random.Random(SEED)
    eval_male = sorted(rng.sample(male_ids, EVAL_MALE_TARGET))
    eval_female = sorted(rng.sample(female_ids, EVAL_FEMALE_TARGET))
    eval_male_set, eval_female_set = set(eval_male), set(eval_female)
    eval_all_set = eval_male_set | eval_female_set

    train_male = sorted(set(male_ids) - eval_male_set)
    train_female = sorted(set(female_ids) - eval_female_set)

    # ---- NEW in v4: per-speaker meeting split (enroll meetings vs test meetings) ----
    splits, skipped_1meeting = compute_meeting_splits(ihm_rows, eval_all_set)
    enroll_meeting_lookup = {sid: set(v["enroll_meetings"]) for sid, v in splits.items()}
    test_meeting_lookup = {sid: set(v["test_meetings"]) for sid, v in splits.items()}

    columns = ["wav_path", "speaker_id", "gender", "location", "language",
               "participant_number", "role_suffix", "native_split",
               "meeting_id", "mic_type"]

    def to_row(r):
        dec = decode_map[r["speaker_id"]]
        return {
            "wav_path": r["wav_path"], "speaker_id": r["speaker_id"],
            "gender": dec["gender"], "location": dec["location"], "language": dec["language"],
            "participant_number": dec["number"], "role_suffix": dec["role_suffix"] or "",
            "native_split": r["native_split"], "meeting_id": r["meeting_id"] or "",
            "mic_type": r["mic_type"],
        }

    enroll_rows = [
        to_row(r) for r in ihm_rows
        if r["speaker_id"] in enroll_meeting_lookup
        and r["meeting_id"] in enroll_meeting_lookup[r["speaker_id"]]
    ]
    test_rows = [
        to_row(r) for r in all_rows
        if r["speaker_id"] in test_meeting_lookup
        and r["meeting_id"] in test_meeting_lookup[r["speaker_id"]]
    ]

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx_path = out_dir / "eval_enroll_test.xlsx"

    wb = Workbook()
    wb.remove(wb.active)

    write_data_sheet(wb, "Enroll", enroll_rows, columns)
    write_data_sheet(wb, "Test", test_rows, columns)

    # ---- NEW in v4: Meeting_Split_Detail audit sheet ----
    split_cols = ["speaker_id", "gender", "n_meetings", "enroll_meetings", "test_meetings",
                  "enroll_row_count", "test_row_count"]
    enroll_count_by_speaker = defaultdict(int)
    for r in enroll_rows:
        enroll_count_by_speaker[r["speaker_id"]] += 1
    test_count_by_speaker = defaultdict(int)
    for r in test_rows:
        test_count_by_speaker[r["speaker_id"]] += 1

    split_rows = []
    for sid in sorted(splits):
        v = splits[sid]
        split_rows.append({
            "speaker_id": sid,
            "gender": decode_map[sid]["gender"],
            "n_meetings": v["n_meetings"],
            "enroll_meetings": ", ".join(v["enroll_meetings"]),
            "test_meetings": ", ".join(v["test_meetings"]),
            "enroll_row_count": enroll_count_by_speaker.get(sid, 0),
            "test_row_count": test_count_by_speaker.get(sid, 0),
        })
    write_data_sheet(wb, "Meeting_Split_Detail", split_rows, split_cols)

    # ---- NEW in v4: Dropped_Speakers sheet (the <2-meeting eval speakers) ----
    dropped_cols = ["speaker_id", "gender", "location", "language", "reason"]
    dropped_rows = []
    for sid in sorted(skipped_1meeting):
        dec = decode_map[sid]
        dropped_rows.append({
            "speaker_id": sid, "gender": dec["gender"], "location": dec["location"],
            "language": dec["language"],
            "reason": "fewer than 2 distinct meetings -- can't split into a separate enroll meeting and test meeting",
        })
    write_data_sheet(wb, "Dropped_Speakers", dropped_rows, dropped_cols)

    train_cols = ["speaker_id", "gender", "location", "language", "participant_number", "role_suffix"]
    train_rows = []
    for sid in sorted(train_male + train_female):
        dec = decode_map[sid]
        train_rows.append({
            "speaker_id": sid, "gender": dec["gender"], "location": dec["location"],
            "language": dec["language"], "participant_number": dec["number"],
            "role_suffix": dec["role_suffix"] or "",
        })
    write_data_sheet(wb, "Train_Group_Speakers", train_rows, train_cols)

    # ---- Summary sheet ----
    ws = wb.create_sheet("Summary", 0)
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 16
    bold = Font(name=FONT_NAME, bold=True)
    normal = Font(name=FONT_NAME)
    n_enroll = len(enroll_rows) + 1  # +1 for header row offset
    n_test = len(test_rows) + 1

    rows_to_write = [
        ("AMI eval group — enroll/test summary (v4: meeting-split scheme)", None, True),
        ("", None, False),
        ("Eval group size (target)", 100, False),
        ("  eval male speakers", EVAL_MALE_TARGET, False),
        ("  eval female speakers", EVAL_FEMALE_TARGET, False),
        ("  of which usable for Enroll/Test (>=2 meetings)", len(splits), False),
        ("  of which DROPPED (only 1 meeting -- see Dropped_Speakers sheet)", len(skipped_1meeting), False),
        ("Train group size (deferred)", len(train_male) + len(train_female), False),
        ("  train male speakers", len(train_male), False),
        ("  train female speakers", len(train_female), False),
        ("Random seed used", SEED, False),
        ("", None, False),
        ("Enroll rows (Enroll sheet)", f"=COUNTA(Enroll!A2:A{max(n_enroll,2)})", False),
        ("  enroll rows, mic_type=ihm (sanity check, should equal total)",
         f"=COUNTIFS(Enroll!J2:J{max(n_enroll,2)},\"ihm\")", False),
        ("Test rows (Test sheet)", f"=COUNTA(Test!A2:A{max(n_test,2)})", False),
        ("  test rows, mic_type=ihm", f"=COUNTIFS(Test!J2:J{max(n_test,2)},\"ihm\")", False),
        ("  test rows, mic_type=sdm", f"=COUNTIFS(Test!J2:J{max(n_test,2)},\"sdm\")", False),
        ("", None, False),
        ("Speaker decode failures (not classified, excluded everywhere)",
         "; ".join(decode_failures) if decode_failures else "none", False),
        ("", None, False),
        ("Notes", None, True),
        ("v4 CHANGE: for each eligible eval speaker, their distinct meetings are sorted and cut in", None, False),
        ("half. The first floor(n/2) meetings go to Enroll (ihm only); the rest go to Test (ihm+sdm).", None, False),
        ("Enroll and Test NEVER share a meeting for the same speaker now -- this is guaranteed by", None, False),
        ("construction, not filtered afterward. Odd meeting counts give Enroll the smaller half.", None, False),
        ("Speakers with only 1 distinct meeting are skipped entirely (can't be split) -- see the", None, False),
        ("Dropped_Speakers sheet. See Meeting_Split_Detail for exactly which meetings went where,", None, False),
        ("per speaker, so any single speaker's split can be checked by hand.", None, False),
        ("No .wav/.txt files were moved, renamed, or modified — this workbook only records paths.", None, False),
    ]
    for i, (label, value, is_bold) in enumerate(rows_to_write, start=1):
        c1 = ws.cell(row=i, column=1, value=label)
        c1.font = bold if is_bold else normal
        if value is not None:
            c2 = ws.cell(row=i, column=2, value=value)
            c2.font = normal

    wb.save(xlsx_path)

    summary = {
        "resolved_root": str(resolved_root),
        "decode_failures": decode_failures,
        "eval_male_speakers": len(eval_male),
        "eval_female_speakers": len(eval_female),
        "eligible_speakers_used_in_enroll_test": len(splits),
        "dropped_speakers_only_1_meeting": skipped_1meeting,
        "train_male_speakers": len(train_male),
        "train_female_speakers": len(train_female),
        "enroll_rows": len(enroll_rows),
        "test_rows": len(test_rows),
        "test_rows_by_mic": {
            "ihm": sum(1 for r in test_rows if r["mic_type"] == "ihm"),
            "sdm": sum(1 for r in test_rows if r["mic_type"] == "sdm"),
        },
        "output_file": str(xlsx_path.resolve()),
    }
    with open(out_dir / "build_summary.json", "w") as f:
        json.dump(summary, f, indent=2)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()