#!/usr/bin/env python3
"""
ami_exact_trial_count.py (v2 — updated for the meeting-split Enroll/Test
scheme) — computes the EXACT total positive and negative trial-pair
counts (no cap, no sampling, no row generation) under your two
constraints:
  1. gender must match between the enroll row and the test row
  2. meeting_id must differ between the enroll row and the test row

This does NOT generate the multi-billion-row trial list. It only counts
it, using per-(speaker, meeting, mic_type) aggregates — the same trick
used throughout this project to get real numbers out of ~100-190 rows of
aggregate data instead of scanning every one of millions of files.

WHAT CHANGED FROM v1 (important if you're comparing old numbers):
v1 assumed Test = ihm+sdm for ALL of a speaker's meetings, and Enroll's
ihm rows were a SUBSET of Test (same meetings, deliberately overlapping)
— so v1 only needed to read the Test sheet.

The xlsx builder changed since then (per your "some change in
architecture" instruction): now, per eligible speaker, their meetings are
split in half — the first floor(n/2) meetings go to Enroll (ihm only),
the rest go to Test (ihm+sdm). Enroll and Test meetings are DISJOINT for
a given speaker now, not overlapping. So this version reads BOTH the
Enroll sheet and the Test sheet independently (Enroll is no longer
derivable from Test), and it also VERIFIES the disjointness on your real
file rather than assuming it — see "Disjointness check" below.

Speakers with fewer than 2 meetings were dropped entirely by the builder
(9 of the 100 real eval speakers) — they simply won't appear in either
sheet, and this script does not try to guess anything about them.

WHAT IT COMPUTES (per speaker s, over their meetings):
  Positive pairs for s  = Enroll_total(s) * Test_total(s)
                           - sum over any meeting m appearing on BOTH
                             sides of Enroll(s,m) * Test(s,m)
  That subtracted term should now come out to exactly 0 for every real
  speaker (Enroll and Test meetings are disjoint by construction) — the
  script still computes it in general (doesn't assume 0) so it would
  catch a real problem if the xlsx were ever built wrong, and it reports
  whether every speaker actually hit 0 so you have real proof, not just
  a design claim.

  Negative pairs for an unordered same-gender pair (s1, s2), both directions:
    Enroll(s1) x Test(s2)  minus sum over meetings SHARED by s1's enroll
                                  side and s2's test side of
                                  Enroll(s1,m) * Test(s2,m)
    Enroll(s2) x Test(s1)  minus the mirror image
  summed over every same-gender pair of eligible eval speakers. This
  filter still matters for negatives — two DIFFERENT speakers' enroll/test
  meeting assignments can still collide even though a single speaker's
  own enroll/test never can.

No pandas anywhere (this environment has a known pandas/numpy ABI bug —
see ami_way_b_feasibility_check.py). openpyxl + plain dict/Counter only.

Usage:
    python ami_exact_trial_count.py eval_xlsx_out/eval_enroll_test.xlsx
    python ami_exact_trial_count.py eval_xlsx_out/eval_enroll_test.xlsx --out exact_trial_counts.json
    python ami_exact_trial_count.py --selftest      # runs a built-in check against
                                                      # hand-verified toy numbers, no xlsx needed
"""

import argparse
import itertools
import json
import sys
from collections import defaultdict
from pathlib import Path

REQUIRED_COLS = {"speaker_id", "meeting_id", "mic_type", "gender"}


# ----------------------------------------------------------------------
# Reading (openpyxl only — do not touch pandas, see module docstring)
# ----------------------------------------------------------------------

def build_speaker_meeting_mic(wb, sheet_name, required_cols=REQUIRED_COLS):
    """Returns (by_speaker, gender_by_speaker, n_rows) where
    by_speaker: speaker_id -> meeting_id -> {"ihm": n, "sdm": n}"""
    if sheet_name not in wb.sheetnames:
        print(f"ERROR: no '{sheet_name}' sheet found. Sheets present: {wb.sheetnames}", file=sys.stderr)
        sys.exit(1)

    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(rows_iter)]
    missing = required_cols - set(header)
    if missing:
        print(f"ERROR: sheet '{sheet_name}' missing columns {missing}. "
              f"Found columns: {header}", file=sys.stderr)
        sys.exit(1)
    sid_i = header.index("speaker_id")
    mtg_i = header.index("meeting_id")
    mic_i = header.index("mic_type")
    gender_i = header.index("gender") if "gender" in header else None

    by_speaker = defaultdict(lambda: defaultdict(lambda: {"ihm": 0, "sdm": 0}))
    gender_by_speaker = {}

    n_rows = 0
    n_bad_mic = 0
    for row in rows_iter:
        if row is None:
            continue
        sid = row[sid_i]
        mtg = row[mtg_i]
        mic = row[mic_i]
        if sid is None or mtg is None or mic is None:
            continue
        mic_norm = str(mic).strip().lower()
        if mic_norm not in ("ihm", "sdm"):
            n_bad_mic += 1
            continue
        sid = str(sid)
        mtg = str(mtg)
        by_speaker[sid][mtg][mic_norm] += 1
        n_rows += 1
        if gender_i is not None and sid not in gender_by_speaker:
            g = row[gender_i]
            if g is not None:
                gender_by_speaker[sid] = str(g)

    if n_bad_mic:
        print(f"WARNING: {n_bad_mic} rows in '{sheet_name}' had a mic_type "
              f"that wasn't 'ihm' or 'sdm' and were skipped.", file=sys.stderr)

    return by_speaker, gender_by_speaker, n_rows


def decode_gender_from_speaker_id(sid: str):
    if not sid:
        return None
    g = sid[0].lower()
    if g == "f":
        return "Female"
    if g == "m":
        return "Male"
    return None


# ----------------------------------------------------------------------
# Disjointness check (replaces v1's "Enroll subset of Test" cross-check,
# since the new builder makes Enroll and Test meetings disjoint instead)
# ----------------------------------------------------------------------

def check_disjoint(enroll_by_speaker, test_by_speaker):
    all_speakers = sorted(set(enroll_by_speaker) | set(test_by_speaker))
    only_enroll = [s for s in all_speakers if s not in test_by_speaker]
    only_test = [s for s in all_speakers if s not in enroll_by_speaker]
    overlap_violations = []
    for sid in all_speakers:
        e_meetings = set(enroll_by_speaker.get(sid, {}))
        t_meetings = set(test_by_speaker.get(sid, {}))
        shared = e_meetings & t_meetings
        if shared:
            overlap_violations.append({"speaker_id": sid, "shared_meetings": sorted(shared)})
    return {
        "n_speakers_in_enroll": len(enroll_by_speaker),
        "n_speakers_in_test": len(test_by_speaker),
        "speakers_only_in_enroll": only_enroll,
        "speakers_only_in_test": only_test,
        "speakers_with_enroll_test_meeting_overlap": len(overlap_violations),
        "overlap_detail": overlap_violations[:50],
        "fully_disjoint_for_every_speaker": len(overlap_violations) == 0,
    }


# ----------------------------------------------------------------------
# Exact positive-pair count
# ----------------------------------------------------------------------

def compute_positive_pairs(enroll_by_speaker, test_by_speaker):
    """enroll_by_speaker / test_by_speaker: speaker_id -> meeting_id ->
    {"ihm": n, "sdm": n}. Enroll should be ihm-only (any "sdm" entries in
    it are flagged, not silently used). Does NOT assume Enroll and Test
    meetings are disjoint -- it computes the same-meeting exclusion in
    general, so it stays correct even if that assumption is ever wrong."""
    per_speaker = {}
    total = 0
    total_ihm_to_ihm = 0
    total_ihm_to_sdm = 0
    speakers_with_nonzero_exclusion = []

    all_speakers = sorted(set(enroll_by_speaker) | set(test_by_speaker))
    for sid in all_speakers:
        e_meetings = enroll_by_speaker.get(sid, {})
        t_meetings = test_by_speaker.get(sid, {})

        enroll_total = sum(v["ihm"] for v in e_meetings.values())
        test_total = sum(v["ihm"] + v["sdm"] for v in t_meetings.values())
        raw = enroll_total * test_total

        shared = set(e_meetings) & set(t_meetings)
        same_meeting_excl = sum(
            e_meetings[m]["ihm"] * (t_meetings[m]["ihm"] + t_meetings[m]["sdm"])
            for m in shared
        )
        positives = raw - same_meeting_excl
        if same_meeting_excl > 0:
            speakers_with_nonzero_exclusion.append(sid)

        test_ihm_total = sum(v["ihm"] for v in t_meetings.values())
        test_sdm_total = sum(v["sdm"] for v in t_meetings.values())
        excl_ihm_ihm = sum(e_meetings[m]["ihm"] * t_meetings[m]["ihm"] for m in shared)
        excl_ihm_sdm = sum(e_meetings[m]["ihm"] * t_meetings[m]["sdm"] for m in shared)
        comp_ihm_ihm = enroll_total * test_ihm_total - excl_ihm_ihm
        comp_ihm_sdm = enroll_total * test_sdm_total - excl_ihm_sdm

        # NEW: full per-meeting breakdown so the positive_pairs number can be
        # recomputed by hand straight from this JSON, without opening the xlsx.
        # "side" tells you whether that meeting was assigned to Enroll or Test
        # for this speaker (this IS point 3 you asked for -- the enroll/test
        # meeting split, per speaker). If a meeting_id ever showed up on BOTH
        # sides (it shouldn't -- see same_meeting_exclusion above) it's marked
        # "both" rather than silently picking one, so nothing gets hidden.
        all_meetings_for_speaker = sorted(set(e_meetings) | set(t_meetings))
        meeting_utterance_counts = {}
        for m in all_meetings_for_speaker:
            on_enroll = m in e_meetings
            on_test = m in t_meetings
            ihm_n = e_meetings.get(m, {}).get("ihm", 0) + t_meetings.get(m, {}).get("ihm", 0)
            sdm_n = e_meetings.get(m, {}).get("sdm", 0) + t_meetings.get(m, {}).get("sdm", 0)
            side = "both (overlap!)" if (on_enroll and on_test) else ("enroll" if on_enroll else "test")
            meeting_utterance_counts[m] = {
                "ihm": ihm_n,
                "sdm": sdm_n,
                "total_utterances": ihm_n + sdm_n,
                "side": side,
            }

        per_speaker[sid] = {
            "n_enroll_meetings": len(e_meetings),
            "n_test_meetings": len(t_meetings),
            "enroll_total_ihm": enroll_total,
            "test_total_ihm_plus_sdm": test_total,
            "same_meeting_exclusion": same_meeting_excl,
            "positive_pairs": positives,
            "component_ihm_enroll_to_ihm_test": comp_ihm_ihm,
            "component_ihm_enroll_to_sdm_test": comp_ihm_sdm,
            # --- added fields ---
            # 1. that speaker's unique meetings
            "unique_meetings": all_meetings_for_speaker,
            # 3. split of the meetings into enroll vs test, as plain lists
            "enroll_meetings": sorted(e_meetings.keys()),
            "test_meetings": sorted(t_meetings.keys()),
            # 2. utterances in each unique meeting (ihm/sdm/total), plus which
            #    side that meeting landed on -- combines 2 and 3 in one place
            #    so you can recompute enroll_total_ihm and test_total_ihm_plus_sdm
            #    by hand: sum "ihm" over side=="enroll" meetings == enroll_total_ihm;
            #    sum "total_utterances" over side=="test" meetings == test_total_ihm_plus_sdm.
            "meeting_utterance_counts": meeting_utterance_counts,
        }
        total += positives
        total_ihm_to_ihm += comp_ihm_ihm
        total_ihm_to_sdm += comp_ihm_sdm

    return {
        "total_positive_pairs": total,
        "total_component_ihm_to_ihm": total_ihm_to_ihm,
        "total_component_ihm_to_sdm": total_ihm_to_sdm,
        "speakers_with_nonzero_same_meeting_exclusion": speakers_with_nonzero_exclusion,
        "per_speaker": per_speaker,
    }


# ----------------------------------------------------------------------
# Exact negative-pair count
# ----------------------------------------------------------------------

def directional_negative(enroll_meetings, test_meetings):
    """enroll_meetings: meeting_id -> ihm_count (this speaker's enroll side).
    test_meetings: meeting_id -> ihm+sdm_count (the other speaker's test side).
    Returns (raw, excluded_by_shared_meeting, surviving)."""
    e_total = sum(enroll_meetings.values())
    t_total = sum(test_meetings.values())
    raw = e_total * t_total
    shared = set(enroll_meetings) & set(test_meetings)
    excl = sum(enroll_meetings[m] * test_meetings[m] for m in shared)
    return raw, excl, raw - excl


def compute_negative_pairs(enroll_by_speaker, test_by_speaker, gender_by_speaker):
    enroll_pm = {sid: {m: v["ihm"] for m, v in meetings.items()}
                 for sid, meetings in enroll_by_speaker.items()}
    test_pm = {sid: {m: v["ihm"] + v["sdm"] for m, v in meetings.items()}
               for sid, meetings in test_by_speaker.items()}

    all_speakers = sorted(set(enroll_by_speaker) | set(test_by_speaker))
    speakers_with_gender = [s for s in all_speakers if gender_by_speaker.get(s)]
    missing_gender = [s for s in all_speakers if not gender_by_speaker.get(s)]
    if missing_gender:
        print(f"WARNING: {len(missing_gender)} speakers had no usable gender value "
              f"and are EXCLUDED from the negative-pair count: {missing_gender}",
              file=sys.stderr)

    by_gender = defaultdict(list)
    for s in speakers_with_gender:
        by_gender[gender_by_speaker[s]].append(s)

    total_negative = 0
    total_raw_no_filter = 0
    total_excluded = 0
    per_gender_totals = {}
    teammate_pairs = 0
    stranger_pairs = 0
    pair_details = []  # only kept for pairs with nonzero exclusion, to keep JSON small

    for gender, speakers in by_gender.items():
        g_total = 0
        g_raw = 0
        for s1, s2 in itertools.combinations(sorted(speakers), 2):
            r12, x12, s12 = directional_negative(enroll_pm.get(s1, {}), test_pm.get(s2, {}))
            r21, x21, s21 = directional_negative(enroll_pm.get(s2, {}), test_pm.get(s1, {}))
            pair_total = s12 + s21
            pair_raw = r12 + r21
            pair_excl = x12 + x21
            g_total += pair_total
            g_raw += pair_raw
            if pair_excl > 0:
                teammate_pairs += 1
                pair_details.append({
                    "speaker_1": s1, "speaker_2": s2,
                    "raw": pair_raw, "excluded_shared_meeting": pair_excl,
                    "surviving": pair_total,
                })
            else:
                stranger_pairs += 1
        per_gender_totals[gender] = {
            "n_speakers": len(speakers),
            "n_pairs": len(speakers) * (len(speakers) - 1) // 2,
            "raw_no_meeting_filter": g_raw,
            "surviving_negative_pairs": g_total,
            "excluded_by_meeting_overlap": g_raw - g_total,
        }
        total_negative += g_total
        total_raw_no_filter += g_raw
        total_excluded += (g_raw - g_total)

    return {
        "total_negative_pairs": total_negative,
        "total_raw_no_meeting_filter": total_raw_no_filter,
        "total_excluded_by_meeting_overlap": total_excluded,
        "by_gender": per_gender_totals,
        "teammate_pair_count": teammate_pairs,
        "stranger_pair_count": stranger_pairs,
        "teammate_pairs_detail": pair_details,
    }


# ----------------------------------------------------------------------
# Self-test against hand-verified toy numbers (v2 scheme: disjoint
# enroll/test meetings per speaker)
# ----------------------------------------------------------------------

def run_selftest():
    # S1: enroll meeting A (150 ihm), test meeting B (150 ihm + 150 sdm)
    #     -- disjoint by construction, exclusion should be 0.
    #     raw = 150 * 300 = 45000, excl = 0 -> positives = 45000
    # S2: shares meeting B on its ENROLL side with S1's TEST side (a
    #     realistic "different speaker overlap" case) -- enroll meeting B
    #     (150 ihm), test meeting C (150 ihm + 150 sdm).
    # S3: enroll meeting D (150 ihm), test meeting E (150 ihm+150 sdm) --
    #     a "stranger" to S1: shares nothing with S1 at all.
    enroll_by_speaker = {
        "mie001": {"A": {"ihm": 150, "sdm": 0}},
        "mie002": {"B": {"ihm": 150, "sdm": 0}},
        "mie003": {"D": {"ihm": 150, "sdm": 0}},
    }
    test_by_speaker = {
        "mie001": {"B": {"ihm": 150, "sdm": 150}},
        "mie002": {"C": {"ihm": 150, "sdm": 150}},
        "mie003": {"E": {"ihm": 150, "sdm": 150}},
    }
    gender_by_speaker = {"mie001": "Male", "mie002": "Male", "mie003": "Male"}

    pos = compute_positive_pairs(enroll_by_speaker, test_by_speaker)
    neg = compute_negative_pairs(enroll_by_speaker, test_by_speaker, gender_by_speaker)

    checks = []
    # Every speaker's own enroll/test meetings are disjoint -> 0 exclusion, full raw survives
    checks.append(("S1 positive pairs == 45000 (150*300, no exclusion)",
                   pos["per_speaker"]["mie001"]["positive_pairs"] == 45000))
    checks.append(("S1 same_meeting_exclusion == 0",
                   pos["per_speaker"]["mie001"]["same_meeting_exclusion"] == 0))
    checks.append(("No speaker has nonzero positive exclusion",
                   pos["speakers_with_nonzero_same_meeting_exclusion"] == []))
    checks.append(("total positives == 45000*3", pos["total_positive_pairs"] == 45000 * 3))

    # new fields: unique_meetings / enroll_meetings / test_meetings / meeting_utterance_counts
    s1 = pos["per_speaker"]["mie001"]
    checks.append(("S1 unique_meetings == ['A','B']", s1["unique_meetings"] == ["A", "B"]))
    checks.append(("S1 enroll_meetings == ['A']", s1["enroll_meetings"] == ["A"]))
    checks.append(("S1 test_meetings == ['B']", s1["test_meetings"] == ["B"]))
    checks.append(("S1 meeting A detail correct (enroll, 150 ihm, 0 sdm)",
                   s1["meeting_utterance_counts"]["A"] == {"ihm": 150, "sdm": 0, "total_utterances": 150, "side": "enroll"}))
    checks.append(("S1 meeting B detail correct (test, 150 ihm, 150 sdm)",
                   s1["meeting_utterance_counts"]["B"] == {"ihm": 150, "sdm": 150, "total_utterances": 300, "side": "test"}))
    # hand-recompute positive_pairs from the new per-meeting fields alone
    recomputed_enroll_total = sum(v["ihm"] for m, v in s1["meeting_utterance_counts"].items() if v["side"] == "enroll")
    recomputed_test_total = sum(v["total_utterances"] for m, v in s1["meeting_utterance_counts"].items() if v["side"] == "test")
    checks.append(("recompute positive_pairs by hand from meeting_utterance_counts alone matches",
                   recomputed_enroll_total * recomputed_test_total == s1["positive_pairs"]))

    # Negative: mie001 enroll(A) -> mie002 test(C): raw=150*300=45000, no shared meeting -> 45000
    #           mie002 enroll(B) -> mie001 test(B): SHARED meeting B! raw=150*300=45000, excl=150*300=45000 -> 0
    # so pair (mie001, mie002) total = 45000 + 0 = 45000, with exclusion 45000
    checks.append(("mie001<->mie002 pair has nonzero exclusion (real overlap case)",
                   any(p["speaker_1"] == "mie001" and p["speaker_2"] == "mie002" for p in neg["teammate_pairs_detail"])))
    checks.append(("teammate_pair_count == 1 (only mie001<->mie002 overlaps)",
                   neg["teammate_pair_count"] == 1))
    checks.append(("stranger_pair_count == 2", neg["stranger_pair_count"] == 2))

    ok = True
    for name, passed in checks:
        print(f"  [{'PASS' if passed else 'FAIL'}] {name}")
        ok = ok and passed
    print("\nSELFTEST", "PASSED" if ok else "FAILED")
    sys.exit(0 if ok else 1)


# ----------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", help="Path to eval_enroll_test.xlsx")
    ap.add_argument("--out", default="exact_trial_counts.json")
    ap.add_argument("--selftest", action="store_true",
                     help="Run built-in correctness check against hand-verified toy numbers, no xlsx needed")
    args = ap.parse_args()

    if args.selftest:
        run_selftest()
        return

    if not args.xlsx:
        print("ERROR: provide the path to eval_enroll_test.xlsx (or run --selftest).", file=sys.stderr)
        sys.exit(1)

    xlsx_path = Path(args.xlsx)
    if not xlsx_path.is_file():
        print(f"ERROR: {xlsx_path} not found.", file=sys.stderr)
        sys.exit(1)

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    enroll_by_speaker, gender_enroll, n_enroll_rows = build_speaker_meeting_mic(wb, "Enroll")
    test_by_speaker, gender_test, n_test_rows = build_speaker_meeting_mic(wb, "Test")
    wb.close()

    # flag any stray sdm rows that ended up in Enroll (should be none)
    enroll_sdm_rows = sum(v["sdm"] for meetings in enroll_by_speaker.values() for v in meetings.values())
    if enroll_sdm_rows:
        print(f"WARNING: Enroll sheet contains {enroll_sdm_rows} sdm rows -- "
              f"it's supposed to be ihm-only.", file=sys.stderr)

    gender_by_speaker = dict(gender_test)
    for sid, g in gender_enroll.items():
        gender_by_speaker.setdefault(sid, g)
    for sid in set(enroll_by_speaker) | set(test_by_speaker):
        if sid not in gender_by_speaker:
            g = decode_gender_from_speaker_id(sid)
            if g:
                gender_by_speaker[sid] = g

    disjoint_check = check_disjoint(enroll_by_speaker, test_by_speaker)

    positives = compute_positive_pairs(enroll_by_speaker, test_by_speaker)
    negatives = compute_negative_pairs(enroll_by_speaker, test_by_speaker, gender_by_speaker)

    grand_total = positives["total_positive_pairs"] + negatives["total_negative_pairs"]

    summary = {
        "input_xlsx": str(xlsx_path.resolve()),
        "enroll_rows_read": n_enroll_rows,
        "test_rows_read": n_test_rows,
        "n_speakers": len(set(enroll_by_speaker) | set(test_by_speaker)),
        "disjointness_check": {
            "fully_disjoint_for_every_speaker": disjoint_check["fully_disjoint_for_every_speaker"],
            "speakers_with_enroll_test_meeting_overlap": disjoint_check["speakers_with_enroll_test_meeting_overlap"],
            "speakers_only_in_enroll": disjoint_check["speakers_only_in_enroll"],
            "speakers_only_in_test": disjoint_check["speakers_only_in_test"],
        },
        "positive_pairs": {
            "total": positives["total_positive_pairs"],
            "component_ihm_enroll_to_ihm_test": positives["total_component_ihm_to_ihm"],
            "component_ihm_enroll_to_sdm_test": positives["total_component_ihm_to_sdm"],
            "speakers_with_nonzero_same_meeting_exclusion": positives["speakers_with_nonzero_same_meeting_exclusion"],
        },
        "negative_pairs": {
            "total": negatives["total_negative_pairs"],
            "raw_no_meeting_filter": negatives["total_raw_no_meeting_filter"],
            "excluded_by_meeting_overlap": negatives["total_excluded_by_meeting_overlap"],
            "by_gender": negatives["by_gender"],
            "teammate_pair_count": negatives["teammate_pair_count"],
            "stranger_pair_count": negatives["stranger_pair_count"],
        },
        "grand_total_trial_pairs": grand_total,
        "positive_to_negative_ratio": (
            f"1 : {round(negatives['total_negative_pairs'] / positives['total_positive_pairs'], 1)}"
            if positives["total_positive_pairs"] else None
        ),
    }

    full_output = {
        "summary": summary,
        "positive_pairs_per_speaker": positives["per_speaker"],
        "negative_pairs_teammate_detail": negatives["teammate_pairs_detail"],
        "disjointness_check_detail": disjoint_check,
    }

    with open(args.out, "w") as f:
        json.dump(full_output, f, indent=2)

    print(json.dumps(summary, indent=2))
    print(f"\nFull per-speaker / per-pair detail written to {args.out}")


if __name__ == "__main__":
    main()